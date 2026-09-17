#!/usr/bin/env python3
"""Generate KidneyUS per-manufacturer test subgroup lists for analysis S1.

S1 stratifies the
KidneyUS benchmark test set by scanner manufacturer. The subgroup membership
files it consumes are:

    ${ANNOTATION_DIR}/KidneyUS/test_mfr_<manufacturer>_list.txt

**Provenance / why this script exists.** Manufacturer is NOT in the per-image
`reviewed_labels_*.csv` (those carry only Quality / View / Comments). The
authoritative source is a separate spreadsheet shipped inside the raw KidneyUS
download:

    kidneyUS/OpenKidneyUltrasoundDataSet_TransducerInfo.xlsx
        columns: Filename, Manufacturer, ManufacturerModelName, ...

Historically the `test_mfr_*` lists were hand-built outside the codebase, so
when the train/val/test split was regenerated (dedup-first → `..._v2`) the
manufacturer lists were *intersected* from the old lists rather than
re-derived. That silently dropped every test sample newly promoted into the
split (46 samples → 376/422 coverage). This script re-derives the lists
directly from the spreadsheet so coverage is always exactly the test set, and
so the step is reproducible for anyone using the codebase.

The xlsx covers all 514 KidneyUS samples, so coverage of any test split is
total by construction; the script asserts this.

Run from `src/`:
  uv run python -m data.utils.gen_kidneyus_mfr_lists \\
    --annotation-dir $WORK/Dataset/Ultrasound/Annotations/38_pt_8_bm_7_ext_v2 \\
    --transducer-xlsx $WORK/Dataset/Ultrasound/Raw/KidneyUS.zip \\
    [--dry-run]

`--transducer-xlsx` accepts either the `.xlsx` itself or the raw `KidneyUS.zip`
that contains it (the member is found automatically).
"""

import argparse
import io
import os
import pathlib
import zipfile
from collections import defaultdict

import openpyxl

from data.utils.dataset_split import write_split_file

DATASET = "KidneyUS"
XLSX_MEMBER_SUFFIX = "OpenKidneyUltrasoundDataSet_TransducerInfo.xlsx"

# Map the raw `Manufacturer` strings to the short, stable subgroup slugs used in
# the filenames and the S1 report. Matched on a lowercased substring so minor
# spelling variants (e.g. the truncated "Philips Medical Sys…") still resolve.
# Unknown manufacturers raise — we never want a sample silently dropped again.
_MFR_RULES = [
    ("acuson", "acuson"),
    ("toshiba", "toshiba"),
    ("siemens", "siemens"),
    ("philips", "philips"),
    ("ge healthcare", "ge"),
]


def normalize_manufacturer(raw: str) -> str:
    low = (raw or "").strip().lower()
    for needle, slug in _MFR_RULES:
        if needle in low:
            return slug
    raise ValueError(f"Unrecognized KidneyUS manufacturer string: {raw!r}")


def _xlsx_bytes(path: pathlib.Path) -> bytes:
    """Return the workbook bytes, reading from a .zip member if needed."""
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            members = [n for n in zf.namelist() if n.endswith(XLSX_MEMBER_SUFFIX)]
            if not members:
                raise FileNotFoundError(
                    f"No '*{XLSX_MEMBER_SUFFIX}' member inside {path}"
                )
            return zf.read(members[0])
    return path.read_bytes()


def load_manufacturer_map(xlsx_path: pathlib.Path) -> dict:
    """Filename-stem (no `_anon`, no extension) -> manufacturer slug."""
    wb = openpyxl.load_workbook(io.BytesIO(_xlsx_bytes(xlsx_path)), data_only=True)
    ws = wb["Sheet1"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    fn_i, mfr_i = header.index("Filename"), header.index("Manufacturer")
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        fn = row[fn_i]
        if not fn:
            continue
        out[str(fn).strip()] = normalize_manufacturer(row[mfr_i])
    return out


def sample_to_stem(sample_id: str) -> str:
    """SaUS sample id `100_IM-0246-0030_anon` -> xlsx key `100_IM-0246-0030`."""
    return sample_id[: -len("_anon")] if sample_id.endswith("_anon") else sample_id


def main() -> None:
    p = argparse.ArgumentParser(description="Generate KidneyUS manufacturer subgroup lists (S1).")
    _ann_default = os.environ.get("ANNOTATION_DIR")
    p.add_argument("--annotation-dir", default=_ann_default,
                   required=(_ann_default is None),
                   help="Annotation root (default: $ANNOTATION_DIR). Reads/writes <root>/KidneyUS/.")
    _raw_default = os.environ.get("RAW_DIR")
    p.add_argument("--transducer-xlsx",
                   default=f"{_raw_default}/KidneyUS.zip" if _raw_default else None,
                   required=(_raw_default is None),
                   help="Path to TransducerInfo.xlsx or the raw KidneyUS.zip containing it "
                        "(default: $RAW_DIR/KidneyUS.zip).")
    p.add_argument("--dry-run", action="store_true", help="Report counts without writing files.")
    args = p.parse_args()

    base = pathlib.Path(args.annotation_dir).expanduser() / DATASET
    test_list = base / "test_list.txt"
    if not test_list.is_file():
        raise FileNotFoundError(f"Missing test list: {test_list}")
    test_ids = [l.strip() for l in test_list.read_text().splitlines() if l.strip()]

    mfr_map = load_manufacturer_map(pathlib.Path(args.transducer_xlsx).expanduser())

    groups = defaultdict(list)
    missing = []
    for sid in test_ids:
        slug = mfr_map.get(sample_to_stem(sid))
        if slug is None:
            missing.append(sid)
        else:
            groups[slug].append(sid)

    if missing:
        raise SystemExit(
            f"{len(missing)} test sample(s) absent from the transducer spreadsheet "
            f"(e.g. {missing[:5]}). The xlsx is expected to cover the full dataset; "
            "investigate before writing partial lists."
        )

    total = sum(len(v) for v in groups.values())
    assert total == len(test_ids), f"partition {total} != test {len(test_ids)}"
    print(f"KidneyUS test={len(test_ids)} → manufacturers:")
    for slug in sorted(groups):
        print(f"  {slug:8s} {len(groups[slug]):4d}   ({base / f'test_mfr_{slug}_list.txt'})")

    if args.dry_run:
        print("[dry-run] no files written.")
        return
    for slug, ids in groups.items():
        write_split_file(base / f"test_mfr_{slug}_list.txt", sorted(ids))
    print(f"Wrote {len(groups)} lists covering {total}/{len(test_ids)} test samples.")


if __name__ == "__main__":
    main()
